#!/usr/bin/env python3
"""Merge fictional, inconsistent order CSV exports into a clean Excel workbook."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - exercised by command line
    raise SystemExit("openpyxl is required. Install dependencies with: python3 -m pip install -r requirements.txt") from exc


FIELDS = ("order_date", "order_id", "channel", "product_name", "quantity", "amount")
OUTPUT_HEADERS = ("Order Date", "Order ID", "Channel", "Product", "Original Product", "Quantity", "Amount")
DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%Y.%m.%d", "%Y.%m.%d.", "%Y/%m/%d", "%m/%d/%y", "%Y%m%d")


def parse_date(value: str) -> date:
    """Parse supported export dates and return a real, sortable date."""
    value = value.strip()
    if not value:
        raise ValueError("empty date")
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unsupported date: {value}")


def parse_amount(value: str) -> float:
    """Parse common currency strings, including commas, symbols and 원."""
    text = value.strip()
    if not text:
        raise ValueError("empty amount")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("() ").replace(",", "")
    text = re.sub(r"(?i)(USD|KRW|원|달러|₩|\$|€|£)", "", text).strip()
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        raise ValueError(f"unsupported amount: {value}")
    amount = float(text)
    return -abs(amount) if negative else amount


def normalize_key(value: str) -> str:
    """Case and punctuation insensitive key for headers and product aliases."""
    return re.sub(r"[^a-z0-9가-힣]", "", value.casefold())


def load_mapping(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            header, field = row.get("source_header", "").strip(), row.get("canonical_field", "").strip()
            if not header or not field:
                continue
            if field not in FIELDS:
                raise ValueError(f"Unknown canonical_field {field!r} in {path}")
            mapping[normalize_key(header)] = field
    return mapping


def load_aliases(path: Path) -> dict[str, str]:
    aliases: dict[str, str] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            alias, standard = row.get("raw_alias", "").strip(), row.get("standard_name", "").strip()
            if alias and standard:
                aliases[normalize_key(alias)] = standard
    return aliases


def load_products(path: Path) -> dict[str, str]:
    """Return normalized standard product names from the catalog."""
    products: dict[str, str] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            name = (row.get("standard_name") or "").strip()
            if name:
                products[normalize_key(name)] = name
    return products


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(handle, dialect=dialect)
        headers = reader.fieldnames or []
        return headers, [dict(row) for row in reader]


def map_headers(headers: list[str], mapping: dict[str, str]) -> dict[str, str]:
    found: dict[str, str] = {}
    for header in headers:
        field = mapping.get(normalize_key(header))
        if field:
            if field in found:
                raise ValueError(f"Multiple columns map to {field!r}: {found[field]!r}, {header!r}")
            found[field] = header
    return found


def set_widths(sheet, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width


def style_sheet(sheet, widths: dict[str, int], freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="233B53")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    set_widths(sheet, widths)


def build_workbook(input_dir: Path, output_path: Path, config_dir: Path | None = None) -> dict[str, Any]:
    config_dir = config_dir or Path(__file__).parent / "config"
    mapping = load_mapping(config_dir / "mapping.csv")
    aliases = load_aliases(config_dir / "product_aliases.csv")
    products = load_products(config_dir / "products.csv")
    paths = sorted((path for path in input_dir.glob("*.csv") if path.is_file()),
                   key=lambda path: (path.stem == "monthly_additions", path.name.casefold()))
    if not paths:
        raise ValueError(f"No .csv files found in {input_dir}")

    clean: list[dict[str, Any]] = []
    changes: list[list[Any]] = [["Type", "File", "Source Row", "Details"]]
    unmatched: dict[str, set[str]] = defaultdict(set)
    seen_orders: dict[tuple[str, str, date], tuple[str, int]] = {}
    seen_rows: dict[tuple[str, ...], tuple[str, int]] = {}
    file_counts: list[tuple[str, int, int, int]] = []
    alias_changes = 0
    duplicates = 0
    parse_failures = 0

    for path in paths:
        headers, rows = read_csv(path)
        mapped = map_headers(headers, mapping)
        missing = [field for field in FIELDS if field not in mapped]
        if missing:
            changes.append(["File warning", path.name, "", "Missing mapped fields: " + ", ".join(missing)])
        read_count, blank_count, accepted_count = 0, 0, 0
        for source_row, row in enumerate(rows, start=2):
            read_count += 1
            if not any((value or "").strip() for value in row.values() if value is not None):
                blank_count += 1
                changes.append(["Blank row", path.name, source_row, "Ignored empty row"])
                continue
            raw = {field: (row.get(header) or "").strip() for field, header in mapped.items()}
            order_id = raw.get("order_id", "")
            channel = raw.get("channel", "") or path.stem
            try:
                order_date = parse_date(raw.get("order_date", ""))
                amount = parse_amount(raw.get("amount", ""))
                quantity_text = raw.get("quantity", "1") or "1"
                quantity = float(quantity_text.replace(",", ""))
                if quantity.is_integer():
                    quantity = int(quantity)
            except (ValueError, TypeError) as exc:
                parse_failures += 1
                changes.append(["Parse failure", path.name, source_row, str(exc)])
                continue
            if not order_id:
                parse_failures += 1
                changes.append(["Parse failure", path.name, source_row, "Missing order_id"])
                continue
            normalized_channel = normalize_key(channel)
            order_key = (normalize_key(order_id), normalized_channel, order_date)
            original_product = raw.get("product_name", "")
            row_key = (normalized_channel, order_date.isoformat(), normalize_key(order_id),
                       normalize_key(original_product), str(quantity), str(amount))
            prior = seen_orders.get(order_key) or seen_rows.get(row_key)
            if prior:
                duplicates += 1
                first_file, first_row = prior
                changes.append(["Duplicate removed", path.name, source_row,
                                f"Order {order_id!r} / {channel!r} / {order_date.isoformat()}; "
                                f"kept first read {first_file} row {first_row}"])
                continue
            seen_orders[order_key] = (path.name, source_row)
            seen_rows[row_key] = (path.name, source_row)
            product_key = normalize_key(original_product)
            standard_product = aliases.get(product_key, products.get(product_key, original_product.strip()))
            if original_product and standard_product != original_product:
                alias_changes += 1
                changes.append(["Product alias", path.name, source_row,
                                f"{original_product!r} → {standard_product!r}"])
            if original_product and product_key not in aliases and product_key not in products:
                unmatched[original_product].add(path.name)
            clean.append({"date": order_date, "order_id": order_id, "channel": channel,
                          "product": standard_product, "original_product": original_product,
                          "quantity": quantity, "amount": amount})
            accepted_count += 1
        file_counts.append((path.name, read_count, blank_count, accepted_count))
        changes.append(["File read", path.name, "", f"{read_count} data rows; {blank_count} blank; {accepted_count} retained"])

    clean.sort(key=lambda item: (item["date"], item["channel"], item["order_id"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Clean Data"
    ws.append(OUTPUT_HEADERS)
    for item in clean:
        ws.append([item["date"], item["order_id"], item["channel"], item["product"],
                   item["original_product"], item["quantity"], item["amount"]])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[5].number_format = "#,##0.##"
        row[6].number_format = '$#,##0.00;[Red]-$#,##0.00'
    style_sheet(ws, {"A": 14, "B": 17, "C": 18, "D": 24, "E": 24, "F": 12, "G": 16})

    product_totals: dict[str, list[float]] = defaultdict(lambda: [0, 0.0])
    month_totals: dict[str, list[float]] = defaultdict(lambda: [0, 0.0])
    for item in clean:
        product_totals[item["product"]][0] += 1
        product_totals[item["product"]][1] += item["amount"]
        month = item["date"].strftime("%Y-%m")
        month_totals[month][0] += 1
        month_totals[month][1] += item["amount"]
    for title, headers_out, entries, widths in (
        ("Summary by Product", ("Product", "Order Count", "Total Amount"), sorted(product_totals.items()), {"A": 28, "B": 16, "C": 20}),
        ("Summary by Month", ("Month", "Order Count", "Total Amount"), sorted(month_totals.items()), {"A": 18, "B": 16, "C": 20}),
    ):
        sheet = wb.create_sheet(title)
        sheet.append(headers_out)
        for label, (count, total) in entries:
            sheet.append([label, count, total])
        for row in sheet.iter_rows(min_row=2):
            row[2].number_format = '$#,##0.00;[Red]-$#,##0.00'
        style_sheet(sheet, widths)

    log_sheet = wb.create_sheet("Change Log")
    for entry in changes:
        log_sheet.append(entry)
    style_sheet(log_sheet, {"A": 22, "B": 28, "C": 14, "D": 72})
    unmatched_sheet = wb.create_sheet("Unmatched")
    unmatched_sheet.append(["Unmatched Product", "Source Files"])
    for product, sources in sorted(unmatched.items(), key=lambda pair: pair[0].casefold()):
        unmatched_sheet.append([product, ", ".join(sorted(sources))])
    style_sheet(unmatched_sheet, {"A": 32, "B": 60})
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {"files": file_counts, "rows": len(clean), "duplicates": duplicates,
            "aliases": alias_changes, "parse_failures": parse_failures,
            "total_amount": sum(item["amount"] for item in clean),
            "unmatched_products": len(unmatched), "output": output_path}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_dir", type=Path, help="Directory containing source CSV files")
    parser.add_argument("-o", "--output", type=Path, default=Path("examples/output/merged.xlsx"))
    parser.add_argument("--config-dir", type=Path, help="Directory containing mapping and alias CSV files")
    args = parser.parse_args()
    try:
        stats = build_workbook(args.input_dir, args.output, args.config_dir)
    except (OSError, ValueError, csv.Error) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    print(f"Wrote {stats['output']}: {stats['rows']} clean rows, {stats['duplicates']} duplicates removed, "
          f"{stats['aliases']} product aliases matched, {stats['parse_failures']} parse failures, "
          f"total ${stats['total_amount']:,.2f}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

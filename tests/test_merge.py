import csv
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from merge import build_workbook, load_aliases, load_products, normalize_key, parse_amount, parse_date  # noqa: E402


class ParsingTests(unittest.TestCase):
    def test_supported_date_formats(self):
        expected = date(2026, 9, 3)
        for raw in ("2026-09-03", "09/03/2026", "2026.9.3", "2026/09/03"):
            with self.subTest(raw=raw):
                self.assertEqual(parse_date(raw), expected)

    def test_amount_formats(self):
        for raw in ("$1,200.00", "1200", "1,200원"):
            with self.subTest(raw=raw):
                self.assertEqual(parse_amount(raw), 1200)
        self.assertEqual(parse_amount("(25.50)"), -25.5)

    def test_product_catalog_and_alias_normalization(self):
        aliases = load_aliases(ROOT / "config" / "product_aliases.csv")
        products = load_products(ROOT / "config" / "products.csv")
        self.assertEqual(aliases[normalize_key("Candle (M)")], "Candle - Medium")
        self.assertEqual(products[normalize_key("Candle - Medium")], "Candle - Medium")
        self.assertEqual(products[normalize_key("Ceramic Mug")], "Ceramic Mug")


class MergeTests(unittest.TestCase):
    def test_sample_duplicate_keeps_first_read_original(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "result.xlsx"
            stats = build_workbook(ROOT / "samples", output)
            self.assertEqual(stats["duplicates"], 5)
            wb = load_workbook(output, data_only=True)
            matches = [row for row in wb["Clean Data"].iter_rows(min_row=2, values_only=True)
                       if row[1] == "ON-1001"]
            self.assertEqual(len(matches), 1)
            self.assertEqual(matches[0][0].date(), date(2026, 9, 3))
            self.assertEqual(matches[0][1:3], ("ON-1001", "Online Shop A"))
            self.assertEqual(matches[0][5:], (1, 1200))
            log = " ".join(str(cell.value) for row in wb["Change Log"].iter_rows()
                           for cell in row if cell.value is not None)
            self.assertIn("kept first read online_shop_a.csv", log)

    def test_duplicate_identity_includes_date_and_first_file_wins(self):
        from openpyxl import load_workbook

        headers = ["Order Date", "Order ID", "Channel", "Product Name", "Qty", "Amount"]
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "input"
            folder.mkdir()
            # Alphabetically later source is still processed first; additions are explicitly last.
            for filename, rows in (
                ("online_shop_a.csv", [
                    ["2026-09-03", "X-1", "Online Shop A", "Tea Set", "1", "450"],
                    ["2026-09-04", "X-1", "Online Shop A", "Tea Set", "1", "450"],
                ]),
                ("monthly_additions.csv", [
                    ["2026-09-03", "X-1", "Online Shop A", "Tea Set", "1", "450"],
                ]),
            ):
                with (folder / filename).open("w", newline="", encoding="utf-8") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(headers)
                    writer.writerows(rows)
            output = Path(temp) / "result.xlsx"
            stats = build_workbook(folder, output)
            self.assertEqual(stats["duplicates"], 1)
            wb = load_workbook(output, data_only=True)
            rows = [r for r in wb["Clean Data"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(len(rows), 2)  # Same order and channel on another date is retained.
            self.assertEqual([r[0].date() for r in rows], [date(2026, 9, 3), date(2026, 9, 4)])

    def test_full_run_totals_sheets_and_catalog_matching(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "merged.xlsx"
            stats = build_workbook(ROOT / "samples", output)
            self.assertEqual(stats["rows"], 73)
            self.assertEqual(stats["duplicates"], 5)
            self.assertEqual(stats["parse_failures"], 2)
            self.assertEqual(stats["aliases"], 10)
            wb = load_workbook(output, data_only=True)
            self.assertEqual(wb.sheetnames, ["Clean Data", "Summary by Product", "Summary by Month", "Change Log", "Unmatched"])
            self.assertEqual(wb["Clean Data"].max_row - 1, 73)
            unmatched_names = {wb["Unmatched"].cell(row, 1).value for row in range(2, wb["Unmatched"].max_row + 1)}
            self.assertEqual(unmatched_names, {"Desk Lamp", "Mystery Trinket"})
            clean_products = {wb["Clean Data"].cell(row, 4).value for row in range(2, wb["Clean Data"].max_row + 1)}
            self.assertIn("Ceramic Mug", clean_products)
            product_total = sum(wb["Summary by Product"].cell(row, 3).value for row in range(2, wb["Summary by Product"].max_row + 1))
            month_total = sum(wb["Summary by Month"].cell(row, 3).value for row in range(2, wb["Summary by Month"].max_row + 1))
            self.assertEqual(product_total, 94140)
            self.assertEqual(month_total, 94140)
            months = {wb["Summary by Month"].cell(row, 1).value for row in range(2, wb["Summary by Month"].max_row + 1)}
            self.assertEqual(months, {"2026-09", "2026-10"})
            log_text = " ".join(str(cell.value) for row in wb["Change Log"].iter_rows() for cell in row if cell.value is not None)
            self.assertIn("Duplicate removed", log_text)
            self.assertIn("Parse failure", log_text)


if __name__ == "__main__":
    unittest.main()
